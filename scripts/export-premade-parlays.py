#!/usr/bin/env python3
"""Write the saved premade-parlay snapshot to a spreadsheet and a CSV.

Reads docs/projections/premade-parlays.json. Does not fetch prices.
"""

from __future__ import annotations

import csv
import json
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "projections" / "premade-parlays.json"
XLSX = ROOT / "docs" / "projections" / "sideline-premade-parlays-2026-10-30.xlsx"
CSV = ROOT / "docs" / "projections" / "sideline-premade-parlays-2026-10-30.csv"
ET = ZoneInfo("America/New_York")

NAVY = "1B3A4B"
WHITE = "FFFFFF"

PARLAY_COLUMNS = (
    "Book",
    "League",
    "Type",
    "American price",
    "Event",
    "Start",
    "Legs",
    "Narrative",
    "FanDuel bet count",
)

LEG_COLUMNS = (
    "Book",
    "League",
    "Type",
    "American price",
    "Parlay event",
    "Leg event",
    "Start",
    "Market",
    "Selection",
    "Line",
    "Outcome",
    "Slate",
)


def et_stamp(iso: str | None) -> str:
    if not iso:
        return ""
    moment = datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(ET)
    clock = moment.strftime("%Y-%m-%d %I:%M %p ET")
    return clock.replace(" 0", " ")


def leg_text(leg: dict) -> str:
    if not leg.get("market") and not leg.get("selection"):
        return "leg not attached"
    selection = leg.get("selection") or "selection missing"
    market = leg.get("market") or "market missing"
    line = leg.get("line")
    shown = selection
    if line not in (None, "") and str(line) not in selection:
        shown = f"{selection} {line}"
    return f"{shown} ({market})"


def parlay_rows(snapshot: dict) -> list[dict]:
    rows = []
    for parlay in [*snapshot["draftKings"], *snapshot["fanDuel"]]:
        starts = [leg.get("start") for leg in parlay["legs"] if leg.get("start")]
        event = parlay.get("event") or "; ".join(parlay.get("events") or [])
        rows.append({
            "Book": parlay["book"],
            "League": parlay.get("league") or "",
            "Type": parlay.get("type") or "",
            "American price": parlay.get("american") or "",
            "Event": event,
            "Start": et_stamp(starts[0] if len(set(starts)) == 1 else (starts[0] if len(starts) == 1 else "")),
            "Legs": " + ".join(leg_text(leg) for leg in parlay["legs"]),
            "Narrative": parlay.get("narrative") or "",
            "FanDuel bet count": parlay.get("totalBets") if parlay.get("totalBets") is not None else "",
        })
    return rows


def leg_rows(snapshot: dict) -> list[dict]:
    rows = []
    for parlay in [*snapshot["draftKings"], *snapshot["fanDuel"]]:
        event = parlay.get("event") or "; ".join(parlay.get("events") or [])
        for leg in parlay["legs"]:
            rows.append({
                "Book": parlay["book"],
                "League": parlay.get("league") or "",
                "Type": parlay.get("type") or "",
                "American price": parlay.get("american") or "",
                "Parlay event": event,
                "Leg event": leg.get("event") or "",
                "Start": et_stamp(leg.get("start")),
                "Market": leg.get("market") or "",
                "Selection": leg.get("selection") or "",
                "Line": "" if leg.get("line") is None else leg.get("line"),
                "Outcome": leg.get("outcome") or "",
                "Slate": leg.get("slate") or "",
            })
    return rows


def write_sheet(ws, columns: tuple[str, ...], rows: list[dict]) -> None:
    header = Font(bold=True, color=WHITE, name="Calibri")
    fill = PatternFill("solid", fgColor=NAVY)
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(1, index, name)
        cell.font = header
        cell.fill = fill
    for row_index, row in enumerate(rows, start=2):
        for column_index, name in enumerate(columns, start=1):
            cell = ws.cell(row_index, column_index, row[name])
            cell.font = Font(name="Calibri")
            cell.alignment = Alignment(vertical="center")
            if name in {"American price", "FanDuel bet count"}:
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(2, len(rows) + 1)}"
    for index, name in enumerate(columns, start=1):
        width = 18
        if name in {"Legs", "Narrative", "Event", "Parlay event", "Leg event", "Market", "Selection"}:
            width = 42
        ws.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    xlsx = Path(sys.argv[2]) if len(sys.argv) > 2 else XLSX
    csv_path = Path(sys.argv[3]) if len(sys.argv) > 3 else CSV
    snapshot = json.loads(source.read_text())
    parlays = parlay_rows(snapshot)
    legs = leg_rows(snapshot)
    if len(parlays) != len(snapshot["draftKings"]) + len(snapshot["fanDuel"]):
        raise SystemExit("parlay row count does not match the snapshot")
    steelers = [row for row in parlays if row["American price"] == "+170" and "Steelers" in row["Event"]]
    if not steelers:
        raise SystemExit("Steelers +170 parlay missing from the spreadsheet")

    workbook = Workbook()
    cover = workbook.active
    cover.title = "Read me"
    cover["A1"] = "Premade parlays"
    cover["A1"].font = Font(bold=True, name="Calibri", size=16)
    retrieved = et_stamp(snapshot["retrievedAt"])
    cover["A3"] = f"Retrieved {snapshot['retrievedAt']} UTC ({retrieved})."
    cover["A4"] = "These are the books' combined prices. A priced edge is unlabeled."
    cover["A5"] = "Parlays is one row per card. Legs is one row per leg."
    cover["A6"] = f"FanDuel rows: {len(snapshot['fanDuel'])}. DraftKings rows: {len(snapshot['draftKings'])}."
    cover.column_dimensions["A"].width = 110

    parlays_sheet = workbook.create_sheet("Parlays")
    write_sheet(parlays_sheet, PARLAY_COLUMNS, parlays)
    legs_sheet = workbook.create_sheet("Legs")
    write_sheet(legs_sheet, LEG_COLUMNS, legs)
    workbook.save(xlsx)

    with csv_path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=PARLAY_COLUMNS)
        writer.writeheader()
        writer.writerows(parlays)

    print(f"parlays {len(parlays)} legs {len(legs)}")
    print(xlsx)
    print(csv_path)


if __name__ == "__main__":
    main()
