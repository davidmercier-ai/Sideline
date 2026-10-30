#!/usr/bin/env python3
"""Write the parlay rank to a workbook. Does not fetch prices."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "projections" / "parlay-rank.json"
XLSX = ROOT / "docs" / "projections" / "sideline-parlay-rank-2026-10-30.xlsx"
NAVY = "1B3A4B"
WHITE = "FFFFFF"

RANK_COLUMNS = (
    "rank",
    "rate",
    "score",
    "book",
    "league",
    "american",
    "event",
    "withCount",
    "againstCount",
    "evenCount",
    "ratedCount",
    "unratedCount",
    "legs",
)

PAIR_COLUMNS = ("rank", "higher", "fanDuel", "draftKings", "event", "legs")


def write_sheet(ws, columns, rows):
    header = Font(bold=True, color=WHITE, name="Calibri")
    fill = PatternFill("solid", fgColor=NAVY)
    labels = {
        "rank": "Rank",
        "rate": "Rate",
        "score": "Score",
        "book": "Book",
        "league": "League",
        "american": "American price",
        "event": "Event",
        "withCount": "With",
        "againstCount": "Against",
        "evenCount": "Even",
        "ratedCount": "Rated legs",
        "unratedCount": "Unrated legs",
        "legs": "Legs",
        "higher": "Higher price",
        "fanDuel": "FanDuel",
        "draftKings": "DraftKings",
    }
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(1, index, labels[name])
        cell.font = header
        cell.fill = fill
    for row_index, row in enumerate(rows, start=2):
        for column_index, name in enumerate(columns, start=1):
            cell = ws.cell(row_index, column_index, row.get(name, ""))
            cell.font = Font(name="Calibri")
            cell.alignment = Alignment(vertical="center", wrap_text=(name == "legs"))
            if name == "american":
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(2, len(rows) + 1)}"
    for index, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = 56 if name in {"legs", "event"} else 18


def main():
    payload = json.loads(SOURCE.read_text())
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Read me"
    cover["A1"] = "Parlay rank"
    cover["A1"].font = Font(bold=True, name="Calibri", size=16)
    counts = payload["counts"]
    cover["A3"] = f"Retrieved {payload['retrievedAt']} UTC. Both books."
    cover["A4"] = "Score is agreement with the model vote and the venue-matched number. It is not a win probability and it is not a priced edge."
    cover["A5"] = "A price is ranked only when both books post the same legs and the same numbers. Higher American odds is the better price."
    cover["A6"] = (
        f"With the projection {counts['withProjection']}. Mixed {counts['mixed']}. "
        f"Against {counts['against']}. Unrated {counts['unrated']}."
    )
    cover.column_dimensions["A"].width = 120
    ranked = workbook.create_sheet("Rank")
    write_sheet(ranked, RANK_COLUMNS, payload["rated"])
    prices = workbook.create_sheet("Same card")
    write_sheet(prices, PAIR_COLUMNS, payload["pairs"])
    workbook.save(XLSX)
    print(XLSX)


if __name__ == "__main__":
    main()
