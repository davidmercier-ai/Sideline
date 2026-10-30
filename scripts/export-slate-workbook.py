#!/usr/bin/env python3
"""Turn a slate markdown snapshot into one spreadsheet.

The workbook keeps the operating manual, one row per game, the four models,
the resampling sim, FanDuel and DraftKings prices, and the game-log histories.
It does not refetch prices and it does not invent a number the snapshot left blank.
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SLATE = ROOT / "docs" / "projections" / "2026-10-30-slate.md"
DEFAULT_MANUAL = ROOT / "docs" / "projections" / "MODEL-AND-SIM.md"
DEFAULT_COMMANDMENTS = ROOT / "docs" / "projections" / "COMMANDMENTS.md"
DEFAULT_OUT = ROOT / "docs" / "projections" / "sideline-all-sports-2026-10-30.xlsx"

METHOD_ORDER = (
    "Venue-matched mean",
    "Overall mean",
    "Venue-matched median",
    "Last 3 games",
)

YARD_STATS = (
    ("rushing yards", "Rushing yards"),
    ("gross passing yards (net passing yards + sack yards lost)", "Gross passing yards"),
    ("net passing yards", "Net passing yards"),
    ("sack yards lost", "Sack yards lost"),
    ("total offensive yards (box-score total, not a sum of medians)", "Total offensive yards"),
)

NARRATIVE_PREFIXES = (
    "Convergence is agreement",
    "Synergy means",
    "Outside consensus:",
    "Team totals:",
    "Our win probability:",
    "Priced edge: not supported.",
    "System refinement:",
    "Outside-pick consensus:",
    "Yardage: not applicable",
)

NAVY = "1F4E79"
GREEN = "1E6B4F"
TEAL = "0F6E6E"
PURPLE = "5C4B8A"
AMBER = "8A5A00"
CLAY = "8A3D3D"
SLATE = "334155"
WHITE = "FFFFFF"
ZEBRA = "F4F7FB"


def num(value):
    if value is None:
        return None
    text = str(value).strip()
    if text in {"", "missing", "unavailable", "—", "-", "–"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def split_row(line):
    if not line.startswith("|"):
        return None
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    if cells and set(cells[0].replace("-", "")) == set():
        return None
    if all(re.fullmatch(r":?-{3,}:?", cell or "---") for cell in cells):
        return None
    return cells


def section_name(line):
    title = line.strip().strip("*").strip()
    if title.startswith("1."):
        return "projection"
    if title == "Simulation":
        return "sim"
    if title.startswith("Historical"):
        return "history"
    if title.startswith("2."):
        return "market"
    if title.startswith("3."):
        return "comparison"
    if title.startswith("4.") or title.startswith("Public consensus"):
        return "consensus"
    if title == "Props":
        return "props"
    return None


def parse_side(raw):
    if not raw or raw == "unavailable":
        return None, None
    match = re.fullmatch(r"([+-]?\d+(?:\.\d+)?)(?:\s+([+-]\d+))?", raw.strip())
    if not match:
        return None, None
    price = int(match.group(2)) if match.group(2) else None
    return float(match.group(1)), price


def parse_moneyline(raw):
    if not raw or raw == "unavailable":
        return None, None, None, None
    paired = re.fullmatch(r"([+-]\d+)\s*/\s*([+-]\d+)", raw.strip())
    if paired:
        return None, int(paired.group(1)), None, int(paired.group(2))
    named = re.fullmatch(r"(.+?)\s+([+-]\d+)\s*/\s*(.+?)\s+([+-]\d+)", raw.strip())
    if named:
        return named.group(1), int(named.group(2)), named.group(3), int(named.group(4))
    return None, None, None, None


def parse_total(raw):
    if not raw or raw == "unavailable":
        return None, None, None, None
    match = re.fullmatch(
        r"o([0-9.]+)\s+([+-]\d+)\s*/\s*u([0-9.]+)\s+([+-]\d+)",
        raw.strip(),
    )
    if not match:
        return None, None, None, None
    return float(match.group(1)), int(match.group(2)), float(match.group(3)), int(match.group(4))


def line_value_class(text):
    if not text:
        return None
    if "line-value candidate" in text:
        return "outside-method-range"
    if "no line value" in text:
        return "inside-method-range"
    if "missing-input" in text:
        return "missing-input"
    return None


def sample_counts(note):
    if not note:
        return None, None
    match = re.search(r"for n=(\d+), against n=(\d+)", note)
    if not match:
        return None, None
    return int(match.group(1)), int(match.group(2))


def parse_history_bullet(line):
    missing = re.fullmatch(
        r"- (.+), (overall|home|away): missing game-level history \(n=(\d+)\)",
        line,
    )
    if missing:
        label, split, count = missing.group(1), missing.group(2), int(missing.group(3))
        team, stat = split_team_stat(label)
        return {
            "team": team,
            "stat": stat,
            "split": split,
            "n": count,
            "mean": None,
            "median": None,
            "mode": None,
            "high": None,
            "low": None,
            "note": "missing game-level history",
        }
    match = re.fullmatch(
        r"- (.+), (overall|home|away): n=(\d+), mean ([^,]+), median ([^,]+), mode (.+), high ([^,]+), low (.+)",
        line,
    )
    if not match:
        return None
    team, stat = split_team_stat(match.group(1))
    return {
        "team": team,
        "stat": stat,
        "split": match.group(2),
        "n": int(match.group(3)),
        "mean": num(match.group(4)),
        "median": num(match.group(5)),
        "mode": match.group(6),
        "high": num(match.group(7)),
        "low": num(match.group(8)),
        "note": None,
    }


def split_team_stat(label):
    # "Ravens points for" / "Man Utd goals against" / "Ravens sack yards lost"
    for marker in (
        " points for",
        " points against",
        " goals for",
        " goals against",
        " rushing yards gained",
        " rushing yards allowed",
        " gross passing yards gained",
        " gross passing yards allowed",
        " net passing yards gained",
        " net passing yards allowed",
        " sack yards lost",
        " total offensive yards gained",
        " total offensive yards allowed",
    ):
        if label.endswith(marker):
            return label[: -len(marker)], marker.strip()
    return label, label


def yard_stat_key(text):
    for raw, label in YARD_STATS:
        if raw in text:
            return raw, label
    return None, None


def blank_yard():
    return {
        "away_venue": None,
        "home_venue": None,
        "combined": None,
        "away_overall": None,
        "home_overall": None,
        "away_for_n": None,
        "away_against_n": None,
        "home_for_n": None,
        "home_against_n": None,
        "away_note": None,
        "home_note": None,
    }


def parse_score_bullet(line, away, home):
    match = re.fullmatch(
        r"- (.+) (points|goals): ([^ ]+) \((.+)\)",
        line,
    )
    if not match:
        return None
    team, unit, value, note = match.group(1), match.group(2), match.group(3), match.group(4)
    if team not in {away, home}:
        return None
    for_n, against_n = sample_counts(note)
    return {
        "team": team,
        "unit": unit,
        "value": num(value),
        "note": note,
        "for_n": for_n,
        "against_n": against_n,
        "thin": "thin sample" in note,
    }


def parse_games(text):
    retrieved = None
    retrieved_match = re.search(r"Retrieved ([0-9T:.\-]+Z)", text)
    if retrieved_match:
        retrieved = retrieved_match.group(1)

    lines = text.splitlines()
    sport = None
    games = []
    unparsed = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("## "):
            heading = line[3:].strip()
            if heading not in {"Open gaps in this snapshot", "Player props"}:
                sport = heading.replace(" scoreboard", "")
            else:
                sport = None
            index += 1
            continue
        if line.startswith("### ") and sport:
            block = [line]
            index += 1
            while index < len(lines) and not lines[index].startswith("### ") and not lines[index].startswith("## "):
                block.append(lines[index])
                index += 1
            game, missed = parse_game(block, sport, retrieved, len(games) + 1)
            games.append(game)
            unparsed.extend(missed)
            continue
        index += 1
    return games, unparsed, retrieved


def parse_game(block, sport, retrieved, order):
    title = block[0][4:].strip()
    if " at " not in title:
        raise SystemExit(f"Game title has no ' at ' separator: {title}")
    away, home = title.split(" at ", 1)
    game = {
        "order": order,
        "sport": sport,
        "away": away,
        "home": home,
        "matchup": title,
        "retrieved": retrieved,
        "when": None,
        "site": None,
        "season_source": None,
        "sample_filter": None,
        "scoring_unit": None,
        "methods": {},
        "comparison_methods": {},
        "history": [],
        "yards": {},
        "identity": None,
        "market": {},
        "market_notes": [],
        "props_flag": None,
        "unparsed": [],
    }
    for raw, label in YARD_STATS:
        game["yards"][raw] = blank_yard()
    section = "meta"
    missed = []

    def keep(line):
        missed.append(f"{title}: {line}")

    for line in block[1:]:
        if line.startswith("**"):
            name = section_name(line)
            if name:
                section = name
            continue
        if not line.strip():
            continue
        if line.startswith("|"):
            cells = split_row(line)
            if not cells:
                continue
            consume_table(game, section, cells)
            continue
        if not line.startswith("- "):
            if line.startswith("Statistical source:"):
                game["season_source"] = line.split("for ", 1)[-1].split(".")[0].strip()
            elif line.startswith("Sample filter:"):
                game["sample_filter"] = line[len("Sample filter:") :].strip()
            continue
        if any(line[2:].startswith(prefix) for prefix in NARRATIVE_PREFIXES):
            continue
        if section == "meta":
            if line.startswith("- When:"):
                game["when"] = line.split(":", 1)[1].strip()
            elif line.startswith("- Site:"):
                game["site"] = line.split(":", 1)[1].strip()
            else:
                keep(line)
            continue
        if section in {"projection", "sim"} and consume_projection(game, line, away, home):
            continue
        if section == "history":
            parsed = parse_history_bullet(line)
            if parsed:
                game["history"].append(parsed)
            else:
                keep(line)
            continue
        if section == "market":
            game["market_notes"].append(line[2:])
            continue
        if section == "comparison" and consume_comparison(game, line):
            continue
        if section == "consensus":
            game.setdefault("consensus_notes", []).append(line[2:])
            continue
        if section == "props":
            continue
        keep(line)
    game["unparsed"] = missed
    return game, missed


def consume_table(game, section, cells):
    if section in {"projection", "sim"} and cells[0] in METHOD_ORDER and len(cells) >= 5:
        game["methods"][cells[0]] = {
            "away": num(cells[1]),
            "home": num(cells[2]),
            "margin": num(cells[3]),
            "combined": num(cells[4]),
        }
        return
    if section == "market" and cells[0] in {
        "Side",
        "Moneyline",
        "O/U",
        "Team total, away",
        "Team total, home",
    }:
        game["market"][cells[0]] = {
            "fd_open": cells[1] if len(cells) > 1 else None,
            "fd_latest": cells[2] if len(cells) > 2 else None,
            "dk_open": cells[3] if len(cells) > 3 else None,
            "dk_latest": cells[4] if len(cells) > 4 else None,
            "flag": cells[5] if len(cells) > 5 else None,
        }
        return
    if section == "comparison" and cells[0] in METHOD_ORDER and len(cells) >= 7:
        game["comparison_methods"][cells[0]] = {
            "margin": num(cells[1]),
            "minus_fd": num(cells[2]),
            "minus_dk": num(cells[3]),
            "combined": num(cells[4]),
            "minus_fd_total": num(cells[5]),
            "minus_dk_total": num(cells[6]),
        }
        return
    if section == "comparison" and cells[0] in {"DraftKings latest", "FanDuel latest"}:
        book = "dk" if cells[0].startswith("DraftKings") else "fd"
        game[f"{book}_book_margin"] = num(cells[1]) if len(cells) > 1 else None
        game[f"{book}_book_total"] = num(cells[3]) if len(cells) > 3 else None
        return
    if section == "props" and cells[0].startswith("Player props"):
        game["props_flag"] = cells[-1]


def consume_projection(game, line, away, home):
    score = parse_score_bullet(line, away, home)
    if score:
        side = "away" if score["team"] == away else "home"
        game["scoring_unit"] = score["unit"]
        game[f"{side}_venue"] = score["value"]
        game[f"{side}_for_n"] = score["for_n"]
        game[f"{side}_against_n"] = score["against_n"]
        game[f"{side}_sample_note"] = score["note"]
        game[f"{side}_thin"] = score["thin"]
        return True
    combined = re.fullmatch(r"- Combined (points|goals): (.+)", line)
    if combined:
        game["scoring_unit"] = combined.group(1)
        game["venue_combined"] = num(combined.group(2))
        return True
    margin = re.fullmatch(r"- Home margin: ([^ ]+) \(.*\)", line)
    if margin:
        game["venue_margin"] = num(margin.group(1))
        return True
    outcome = re.fullmatch(r"- Projected outcome: (.+)", line)
    if outcome:
        game["projected_outcome"] = outcome.group(1)
        return True
    overall = re.fullmatch(r"- Overall comparison, (.+): ([^ ]+) \((.+)\)", line)
    if overall and yard_stat_key(overall.group(1))[0] is None:
        team, value, note = overall.group(1), num(overall.group(2)), overall.group(3)
        side = "away" if team == away else "home" if team == home else None
        if side:
            game[f"{side}_overall"] = value
            game[f"{side}_overall_note"] = note
            return True
    convergence = re.fullmatch(
        r"- Convergence: (\d+) methods available; home (\d+), away (\d+), tie (\d+)\. Result: ([^.]+)\.",
        line,
    )
    if convergence:
        game["methods_available"] = int(convergence.group(1))
        game["home_votes"] = int(convergence.group(2))
        game["away_votes"] = int(convergence.group(3))
        game["tie_votes"] = int(convergence.group(4))
        game["agreement"] = convergence.group(5)
        return True
    ranges = re.fullmatch(
        r"- Margin range across methods: ([0-9.]+)\. Combined range across methods: ([0-9.]+)\.",
        line,
    )
    if ranges:
        game["margin_range"] = num(ranges.group(1))
        game["total_range"] = num(ranges.group(2))
        return True
    synergy = re.fullmatch(r"- Synergy, (visitor|home) scoring inputs: (\S+) \((.*)\)\.", line)
    if synergy:
        who = "visitor" if synergy.group(1) == "visitor" else "home"
        game[f"{who}_synergy"] = synergy.group(2)
        game[f"{who}_synergy_detail"] = synergy.group(3)
        return True
    median = re.fullmatch(r"- Sim median score: (.+) ([0-9.]+), (.+) ([0-9.]+)\.", line)
    if median:
        game["sim_away_name"] = median.group(1)
        game["sim_away_median"] = num(median.group(2))
        game["sim_home_name"] = median.group(3)
        game["sim_home_median"] = num(median.group(4))
        return True
    if line.startswith("- Sim frequencies:"):
        body = line.split(":", 1)[1]
        wins = re.findall(r"\s*([^,]+?) wins ([0-9.]+)%", body)
        tie = re.search(r"ties ([0-9.]+)%", body)
        for name, pct in wins:
            if name == away:
                game["sim_away_win_pct"] = num(pct)
            elif name == home:
                game["sim_home_win_pct"] = num(pct)
            else:
                return False
        if tie:
            game["sim_tie_pct"] = num(tie.group(1))
        return True
    sample = re.fullmatch(r"- Sim sample: (.+)", line)
    if sample:
        game["sim_sample"] = sample.group(1)
        return True
    cover = re.fullmatch(
        r"- Sim home-cover rate versus (FanDuel|DraftKings): ([0-9.]+)% \(pushes ([0-9.]+)%\)\.",
        line,
    )
    if cover:
        book = "fd" if cover.group(1) == "FanDuel" else "dk"
        game[f"sim_{book}_cover_pct"] = num(cover.group(2))
        game[f"sim_{book}_push_pct"] = num(cover.group(3))
        return True
    over = re.fullmatch(
        r"- Sim over rate versus (FanDuel|DraftKings) total: ([0-9.]+)%\.",
        line,
    )
    if over:
        book = "fd" if over.group(1) == "FanDuel" else "dk"
        game[f"sim_{book}_over_pct"] = num(over.group(2))
        return True
    if line.startswith("- Price gap:"):
        game["price_gap"] = line.split(":", 1)[1].strip()
        return True
    identity = re.fullmatch(r"- Total-offense identity: (.+)", line)
    if identity:
        game["identity"] = identity.group(1)
        return True
    if consume_yard(game, line, away, home):
        return True
    return False


def consume_yard(game, line, away, home):
    combined = re.fullmatch(r"- Combined (.+): (.+)", line)
    if combined:
        key, _label = yard_stat_key(combined.group(1))
        if key:
            game["yards"][key]["combined"] = num(combined.group(2))
            return True
        return False
    overall = re.fullmatch(r"- Overall comparison, (.+): ([^ ]+) \((.+)\)", line)
    if overall:
        body = overall.group(1)
        key, _label = yard_stat_key(body)
        if not key:
            return False
        team = body[: body.index(key)].strip()
        side = "away" if team == away else "home" if team == home else None
        if not side:
            return False
        game["yards"][key][f"{side}_overall"] = num(overall.group(2))
        return True
    team_line = re.fullmatch(r"- (.+): ([^ ]+) \((.+)\)", line)
    if not team_line:
        return False
    body, value, note = team_line.group(1), team_line.group(2), team_line.group(3)
    key, _label = yard_stat_key(body)
    if not key:
        return False
    team = body[: body.index(key)].strip()
    side = "away" if team == away else "home" if team == home else None
    if not side:
        return False
    for_n, against_n = sample_counts(note)
    game["yards"][key][f"{side}_venue"] = num(value)
    game["yards"][key][f"{side}_for_n"] = for_n
    game["yards"][key][f"{side}_against_n"] = against_n
    game["yards"][key][f"{side}_note"] = note
    return True


def consume_comparison(game, line):
    side_gap = re.fullmatch(
        r"- Venue-mean side gap versus FanDuel: (.+)\. Versus DraftKings: (.+)\.",
        line,
    )
    if side_gap:
        game["side_gap_fd"] = side_gap.group(1)
        game["side_gap_dk"] = side_gap.group(2)
        return True
    total_gap = re.fullmatch(
        r"- Venue-mean total gap versus FanDuel: (.+)\. Versus DraftKings: (.+)\.",
        line,
    )
    if total_gap:
        game["total_gap_fd"] = total_gap.group(1)
        game["total_gap_dk"] = total_gap.group(2)
        return True
    value = re.fullmatch(r"- (FanDuel|DraftKings) (side|total): (.+)", line)
    if value:
        book = "fd" if value.group(1) == "FanDuel" else "dk"
        kind = value.group(2)
        text = value.group(3)
        game[f"{book}_{kind}_value"] = text
        game[f"{book}_{kind}_class"] = line_value_class(text)
        return True
    books = re.fullmatch(
        r"- FanDuel versus DraftKings home margin: ([^ ]+) \(FanDuel minus DraftKings\)\.",
        line,
    )
    if books:
        game["fd_minus_dk_margin"] = num(books.group(1))
        return True
    if line.startswith("- FanDuel versus DraftKings:"):
        game["fd_minus_dk_note"] = line.split(":", 1)[1].strip()
        return True
    novig = re.fullmatch(
        r"- (FanDuel|DraftKings) no-vig price(?:, from the moneyline only)?: (.+)",
        line,
    )
    if novig:
        book = "fd" if novig.group(1) == "FanDuel" else "dk"
        rest = novig.group(2)
        parsed = re.fullmatch(
            r"(.+?) ([0-9.]+)%, (.+?) ([0-9.]+)%, hold (-?[0-9.]+) percentage points\.",
            rest,
        )
        if parsed:
            pairs = ((parsed.group(1), num(parsed.group(2))), (parsed.group(3), num(parsed.group(4))))
            for name, pct in pairs:
                if name == game["away"]:
                    game[f"{book}_novig_away"] = pct
                elif name == game["home"]:
                    game[f"{book}_novig_home"] = pct
            game[f"{book}_hold"] = num(parsed.group(5))
        else:
            game[f"{book}_novig_note"] = rest.rstrip(".")
        return True
    return False


def market_cell(game, market, field):
    return (game["market"].get(market) or {}).get(field)


def game_row(game):
    methods = game["methods"]
    venue = methods.get("Venue-matched mean", {})
    overall = methods.get("Overall mean", {})
    median = methods.get("Venue-matched median", {})
    last3 = methods.get("Last 3 games", {})
    side = game["market"].get("Side", {})
    money = game["market"].get("Moneyline", {})
    total = game["market"].get("O/U", {})
    fd_spread, fd_spread_price = parse_side(side.get("fd_latest"))
    dk_spread, dk_spread_price = parse_side(side.get("dk_latest"))
    dk_open_spread, dk_open_price = parse_side(side.get("dk_open"))
    _fd_ml_away_abbr, fd_ml_away, _fd_ml_home_abbr, fd_ml_home = parse_moneyline(money.get("fd_latest"))
    _dk_ml_away_abbr, dk_ml_away, _dk_ml_home_abbr, dk_ml_home = parse_moneyline(money.get("dk_latest"))
    _dk_open_away_abbr, dk_open_ml_away, _dk_open_home_abbr, dk_open_ml_home = parse_moneyline(money.get("dk_open"))
    fd_total, fd_over, _fd_under_total, fd_under = parse_total(total.get("fd_latest"))
    dk_total, dk_over, _dk_under_total, dk_under = parse_total(total.get("dk_latest"))
    dk_open_total, dk_open_over, _dk_open_under_total, dk_open_under = parse_total(total.get("dk_open"))
    thin = bool(game.get("away_thin") or game.get("home_thin") or (game.get("sim_sample") or "").startswith("thin"))
    hold_note = None
    for book in ("fd", "dk"):
        hold = game.get(f"{book}_hold")
        if hold is not None and hold < 0:
            hold_note = (
                "A negative hold means the two moneyline prices do not form a two-way market. "
                "Soccer prices include a draw this pair leaves out. This is the book’s arithmetic, not our win probability."
            )
    return {
        "order": game["order"],
        "sport": game["sport"],
        "away": game["away"],
        "home": game["home"],
        "when": game["when"],
        "site": game["site"],
        "scoring_unit": game.get("scoring_unit"),
        "season_source": game.get("season_source"),
        "sample_filter": game.get("sample_filter"),
        "model_sample": "thin" if (game.get("away_thin") or game.get("home_thin")) else "at least 3",
        "sim_stability": "thin" if (game.get("sim_sample") or "").startswith("thin") else "stable",
        "thin_any": "yes" if thin else "no",
        "away_for_n": game.get("away_for_n"),
        "away_against_n": game.get("away_against_n"),
        "home_for_n": game.get("home_for_n"),
        "home_against_n": game.get("home_against_n"),
        "venue_away": venue.get("away", game.get("away_venue")),
        "venue_home": venue.get("home", game.get("home_venue")),
        "venue_margin": venue.get("margin", game.get("venue_margin")),
        "venue_combined": venue.get("combined", game.get("venue_combined")),
        "overall_away": overall.get("away", game.get("away_overall")),
        "overall_home": overall.get("home", game.get("home_overall")),
        "overall_margin": overall.get("margin"),
        "overall_combined": overall.get("combined"),
        "median_away": median.get("away"),
        "median_home": median.get("home"),
        "median_margin": median.get("margin"),
        "median_combined": median.get("combined"),
        "last3_away": last3.get("away"),
        "last3_home": last3.get("home"),
        "last3_margin": last3.get("margin"),
        "last3_combined": last3.get("combined"),
        "projected_outcome": game.get("projected_outcome"),
        "methods_available": game.get("methods_available"),
        "home_votes": game.get("home_votes"),
        "away_votes": game.get("away_votes"),
        "tie_votes": game.get("tie_votes"),
        "agreement": game.get("agreement"),
        "margin_range": game.get("margin_range"),
        "total_range": game.get("total_range"),
        "visitor_synergy": game.get("visitor_synergy"),
        "visitor_synergy_detail": game.get("visitor_synergy_detail"),
        "home_synergy": game.get("home_synergy"),
        "home_synergy_detail": game.get("home_synergy_detail"),
        "sim_away_median": game.get("sim_away_median"),
        "sim_home_median": game.get("sim_home_median"),
        "sim_away_win_pct": game.get("sim_away_win_pct"),
        "sim_home_win_pct": game.get("sim_home_win_pct"),
        "sim_tie_pct": game.get("sim_tie_pct"),
        "sim_sample": game.get("sim_sample"),
        "sim_fd_cover_pct": game.get("sim_fd_cover_pct"),
        "sim_fd_push_pct": game.get("sim_fd_push_pct"),
        "sim_dk_cover_pct": game.get("sim_dk_cover_pct"),
        "sim_dk_push_pct": game.get("sim_dk_push_pct"),
        "sim_fd_over_pct": game.get("sim_fd_over_pct"),
        "sim_dk_over_pct": game.get("sim_dk_over_pct"),
        "price_gap": game.get("price_gap"),
        "edge_status": "not labeled",
        "fd_side_open": empty_to_blank(side.get("fd_open")),
        "fd_side_latest": empty_to_blank(side.get("fd_latest")),
        "fd_latest_home_spread": fd_spread,
        "fd_latest_spread_price": fd_spread_price,
        "dk_side_open": empty_to_blank(side.get("dk_open")),
        "dk_side_latest": empty_to_blank(side.get("dk_latest")),
        "dk_open_home_spread": dk_open_spread,
        "dk_open_spread_price": dk_open_price,
        "dk_latest_home_spread": dk_spread,
        "dk_latest_spread_price": dk_spread_price,
        "side_flag": side.get("flag"),
        "fd_ml_open": empty_to_blank(money.get("fd_open")),
        "fd_ml_latest": empty_to_blank(money.get("fd_latest")),
        "fd_ml_away_price": fd_ml_away,
        "fd_ml_home_price": fd_ml_home,
        "dk_ml_open": empty_to_blank(money.get("dk_open")),
        "dk_ml_latest": empty_to_blank(money.get("dk_latest")),
        "dk_open_ml_away": dk_open_ml_away,
        "dk_open_ml_home": dk_open_ml_home,
        "dk_latest_ml_away": dk_ml_away,
        "dk_latest_ml_home": dk_ml_home,
        "ml_flag": money.get("flag"),
        "fd_total_open": empty_to_blank(total.get("fd_open")),
        "fd_total_latest": empty_to_blank(total.get("fd_latest")),
        "fd_latest_total": fd_total,
        "fd_over_price": fd_over,
        "fd_under_price": fd_under,
        "dk_total_open": empty_to_blank(total.get("dk_open")),
        "dk_total_latest": empty_to_blank(total.get("dk_latest")),
        "dk_open_total": dk_open_total,
        "dk_open_over_price": dk_open_over,
        "dk_open_under_price": dk_open_under,
        "dk_latest_total": dk_total,
        "dk_over_price": dk_over,
        "dk_under_price": dk_under,
        "total_flag": total.get("flag"),
        "team_total_away": empty_to_blank(market_cell(game, "Team total, away", "flag")),
        "team_total_home": empty_to_blank(market_cell(game, "Team total, home", "flag")),
        "price_notes": " | ".join(game["market_notes"]) or None,
        "fd_book_margin": game.get("fd_book_margin"),
        "dk_book_margin": game.get("dk_book_margin"),
        "fd_book_total": game.get("fd_book_total"),
        "dk_book_total": game.get("dk_book_total"),
        "fd_side_class": game.get("fd_side_class"),
        "dk_side_class": game.get("dk_side_class"),
        "fd_total_class": game.get("fd_total_class"),
        "dk_total_class": game.get("dk_total_class"),
        "fd_side_value": game.get("fd_side_value"),
        "dk_side_value": game.get("dk_side_value"),
        "fd_total_value": game.get("fd_total_value"),
        "dk_total_value": game.get("dk_total_value"),
        "side_gap_fd": game.get("side_gap_fd"),
        "side_gap_dk": game.get("side_gap_dk"),
        "total_gap_fd": game.get("total_gap_fd"),
        "total_gap_dk": game.get("total_gap_dk"),
        "fd_minus_dk_margin": game.get("fd_minus_dk_margin"),
        "fd_minus_dk_note": game.get("fd_minus_dk_note"),
        "fd_novig_away_pct": game.get("fd_novig_away"),
        "fd_novig_home_pct": game.get("fd_novig_home"),
        "fd_hold_pts": game.get("fd_hold"),
        "dk_novig_away_pct": game.get("dk_novig_away"),
        "dk_novig_home_pct": game.get("dk_novig_home"),
        "dk_hold_pts": game.get("dk_hold"),
        "novig_caution": hold_note,
        "fd_novig_note": game.get("fd_novig_note"),
        "dk_novig_note": game.get("dk_novig_note"),
        "props_flag": game.get("props_flag"),
        "identity": game.get("identity"),
        "snapshot": game.get("retrieved"),
        "sample_filter": game.get("sample_filter"),
    }


def empty_to_blank(value):
    if value in {None, "", "unavailable"}:
        return None
    return value


GAME_COLUMNS = [
    ("order", "Order", 8, SLATE, "Position on the retrieved board."),
    ("sport", "Sport", 28, SLATE, "Board the game came from. Friday and Saturday college are separate."),
    ("away", "Away", 26, SLATE, "Visitor, including a ranking when ESPN printed one."),
    ("home", "Home", 26, SLATE, "Home team. On a neutral site this is still the team listed second."),
    ("when", "When", 24, SLATE, "Kickoff or tip as printed. Times are Eastern. Soccer rows say Scheduled when ESPN sent no clock."),
    ("site", "Site", 32, SLATE, "home team listed second, or neutral (venue split not applied)."),
    ("scoring_unit", "Scoring unit", 14, SLATE, "points for football and NBA. goals for NHL and soccer."),
    ("season_source", "Season source", 36, SLATE, "ESPN season the logs were filtered to."),
    ("model_sample", "Model sample", 14, GREEN, "thin when either venue-matched points or goals sample is under 3 games."),
    ("sim_stability", "Sim stability", 14, PURPLE, "stable when every venue slice has at least 3 games. thin otherwise."),
    ("thin_any", "Thin anywhere", 14, SLATE, "yes when the model sample or the sim sample is thin."),
    ("away_for_n", "Away for n", 12, GREEN, "Visitor venue games in the venue-matched mean."),
    ("away_against_n", "Away against n", 14, GREEN, "Home team’s venue games allowed, used on the visitor score."),
    ("home_for_n", "Home for n", 12, GREEN, "Home team’s venue games in the venue-matched mean."),
    ("home_against_n", "Home against n", 14, GREEN, "Visitor venue games allowed, used on the home score."),
    ("venue_away", "Venue mean away", 16, GREEN, "Model 1. Not blended with the other three."),
    ("venue_home", "Venue mean home", 16, GREEN, "Model 1 home figure."),
    ("venue_margin", "Venue mean margin", 16, GREEN, "Home minus away. Positive means the home team is higher."),
    ("venue_combined", "Venue mean total", 16, GREEN, "Away plus home on model 1."),
    ("overall_away", "Overall mean away", 16, GREEN, "Model 2. All completed games, not the venue slice."),
    ("overall_home", "Overall mean home", 16, GREEN, "Model 2 home figure."),
    ("overall_margin", "Overall mean margin", 18, GREEN, "Model 2 home margin."),
    ("overall_combined", "Overall mean total", 16, GREEN, "Model 2 combined."),
    ("median_away", "Median away", 14, GREEN, "Model 3. Average of the two venue medians. Blank if either median is missing."),
    ("median_home", "Median home", 14, GREEN, "Model 3 home figure."),
    ("median_margin", "Median margin", 14, GREEN, "Model 3 home margin."),
    ("median_combined", "Median total", 14, GREEN, "Model 3 combined."),
    ("last3_away", "Last 3 away", 14, GREEN, "Model 4. Last three completed games, any venue."),
    ("last3_home", "Last 3 home", 14, GREEN, "Model 4 home figure."),
    ("last3_margin", "Last 3 margin", 14, GREEN, "Model 4 home margin."),
    ("last3_combined", "Last 3 total", 14, GREEN, "Model 4 combined."),
    ("projected_outcome", "Venue-mean side", 18, GREEN, "Side named by the venue-matched mean only. A tie is within 0.05."),
    ("methods_available", "Methods voting", 16, TEAL, "How many of the four models had a margin."),
    ("home_votes", "Home votes", 12, TEAL, "Models with a home margin above 0.05."),
    ("away_votes", "Away votes", 12, TEAL, "Models with a home margin below -0.05."),
    ("tie_votes", "Tie votes", 12, TEAL, "Models inside 0.05 of zero."),
    ("agreement", "Agreement", 12, TEAL, "home, away, or split. Split is not a blend and not consensus."),
    ("margin_range", "Margin range", 14, TEAL, "Widest minus narrowest home margin across the four models."),
    ("total_range", "Total range", 12, TEAL, "Widest minus narrowest combined total across the four models."),
    ("visitor_synergy", "Visitor synergy", 16, TEAL, "reinforce, diverge, flat, or missing-input."),
    ("visitor_synergy_detail", "Visitor synergy detail", 42, TEAL, "Own venue rate versus own overall, and opponent venue allowance versus opponent overall."),
    ("home_synergy", "Home synergy", 16, TEAL, "Same test for the home team’s score."),
    ("home_synergy_detail", "Home synergy detail", 42, TEAL, "Home-team input signs."),
    ("sim_away_median", "Sim median away", 16, PURPLE, "Median of 2,000 draws. Seed 20261030 mixed with the ESPN game id."),
    ("sim_home_median", "Sim median home", 16, PURPLE, "Median home score from the same draws."),
    ("sim_away_win_pct", "Sim away win %", 16, PURPLE, "Percent of draws the visitor wins. Already in percent, not a probability scaled to 1."),
    ("sim_home_win_pct", "Sim home win %", 16, PURPLE, "Percent of draws the home team wins. Not a validated win probability."),
    ("sim_tie_pct", "Sim tie %", 12, PURPLE, "Percent of draws with equal scores."),
    ("sim_sample", "Sim sample note", 42, PURPLE, "stable or the thin-sample sentence from the snapshot."),
    ("sim_fd_cover_pct", "Sim FD home cover %", 18, PURPLE, "Draws where the home margin is greater than FanDuel’s home margin."),
    ("sim_fd_push_pct", "Sim FD push %", 14, PURPLE, "Draws equal to the FanDuel home margin."),
    ("sim_dk_cover_pct", "Sim DK home cover %", 18, PURPLE, "Same cover test against DraftKings."),
    ("sim_dk_push_pct", "Sim DK push %", 14, PURPLE, "Draws equal to the DraftKings home margin."),
    ("sim_fd_over_pct", "Sim FD over %", 14, PURPLE, "Draws whose sum is greater than the FanDuel total."),
    ("sim_dk_over_pct", "Sim DK over %", 14, PURPLE, "Draws whose sum is greater than the DraftKings total."),
    ("price_gap", "Price gap note", 42, PURPLE, "Sim home-win frequency minus the nearer no-vig price, in percentage points. Not an edge."),
    ("edge_status", "Priced edge", 14, CLAY, "Always not labeled on this snapshot."),
    ("fd_side_latest", "FD side latest", 16, AMBER, "FanDuel Ohio home spread and price. Negative spread means home is favored. Opening price was not on the board."),
    ("fd_latest_home_spread", "FD home spread", 16, AMBER, "Number only, same sign as the board."),
    ("fd_latest_spread_price", "FD spread price", 16, AMBER, "American odds on that FanDuel spread."),
    ("dk_side_open", "DK side open", 16, AMBER, "DraftKings open spread from the ESPN quote."),
    ("dk_open_home_spread", "DK open spread", 14, AMBER, "Open spread number."),
    ("dk_open_spread_price", "DK open spread price", 18, AMBER, "American odds on the open DraftKings spread."),
    ("dk_side_latest", "DK side latest", 16, AMBER, "DraftKings latest spread from ESPN."),
    ("dk_latest_home_spread", "DK home spread", 16, AMBER, "Latest spread number. Negative means home is favored."),
    ("dk_latest_spread_price", "DK spread price", 16, AMBER, "American odds on the latest DraftKings spread."),
    ("fd_book_margin", "FD home margin", 16, AMBER, "Points by which FanDuel has the home team favored. Equal to the home spread times -1. Blank when the snapshot could not sign it."),
    ("dk_book_margin", "DK home margin", 16, AMBER, "Same sign convention for DraftKings. Blank when the snapshot left it missing."),
    ("fd_ml_latest", "FD moneyline latest", 22, AMBER, "Away price / home price, or unavailable."),
    ("fd_ml_away_price", "FD ML away", 12, AMBER, "American odds."),
    ("fd_ml_home_price", "FD ML home", 12, AMBER, "American odds."),
    ("dk_ml_open", "DK moneyline open", 24, AMBER, "ESPN open moneyline text."),
    ("dk_open_ml_away", "DK open ML away", 16, AMBER, "American odds."),
    ("dk_open_ml_home", "DK open ML home", 16, AMBER, "American odds."),
    ("dk_ml_latest", "DK moneyline latest", 24, AMBER, "ESPN latest moneyline text."),
    ("dk_latest_ml_away", "DK ML away", 12, AMBER, "American odds."),
    ("dk_latest_ml_home", "DK ML home", 12, AMBER, "American odds."),
    ("fd_total_latest", "FD total latest", 22, AMBER, "Over and under text."),
    ("fd_latest_total", "FD total", 12, AMBER, "Posted total."),
    ("fd_over_price", "FD over price", 14, AMBER, "American odds."),
    ("fd_under_price", "FD under price", 14, AMBER, "American odds."),
    ("dk_total_open", "DK total open", 28, AMBER, "ESPN open total text."),
    ("dk_open_total", "DK open total", 14, AMBER, "Open total number."),
    ("dk_open_over_price", "DK open over", 14, AMBER, "American odds."),
    ("dk_open_under_price", "DK open under", 14, AMBER, "American odds."),
    ("dk_total_latest", "DK total latest", 28, AMBER, "ESPN latest total text."),
    ("dk_latest_total", "DK total", 12, AMBER, "Latest total number."),
    ("dk_over_price", "DK over price", 14, AMBER, "American odds."),
    ("dk_under_price", "DK under price", 14, AMBER, "American odds."),
    ("fd_book_total", "FD total used", 14, AMBER, "Total the comparison used. Blank when FanDuel had no total."),
    ("dk_book_total", "DK total used", 14, AMBER, "Total the comparison used."),
    ("price_notes", "Price notes", 42, AMBER, "Same-number ranking, or a note that the numbers differ and were not ranked."),
    ("fd_side_class", "FD side vs models", 22, CLAY, "inside-method-range, outside-method-range, or missing-input. Outside is a point gap, not an edge."),
    ("dk_side_class", "DK side vs models", 22, CLAY, "Same test for DraftKings."),
    ("fd_total_class", "FD total vs models", 22, CLAY, "FanDuel total against the four combined totals."),
    ("dk_total_class", "DK total vs models", 22, CLAY, "DraftKings total against the four combined totals."),
    ("fd_side_value", "FD side note", 42, CLAY, "Sentence from the snapshot."),
    ("dk_side_value", "DK side note", 42, CLAY, "Sentence from the snapshot."),
    ("fd_total_value", "FD total note", 42, CLAY, "Sentence from the snapshot."),
    ("dk_total_value", "DK total note", 42, CLAY, "Sentence from the snapshot."),
    ("side_gap_fd", "Venue margin minus FD", 24, CLAY, "Venue-matched mean margin minus the FanDuel home margin."),
    ("side_gap_dk", "Venue margin minus DK", 24, CLAY, "Venue-matched mean margin minus the DraftKings home margin."),
    ("total_gap_fd", "Venue total minus FD", 36, CLAY, "Includes the above/below wording from the snapshot."),
    ("total_gap_dk", "Venue total minus DK", 36, CLAY, "Includes the above/below wording from the snapshot."),
    ("fd_minus_dk_margin", "FD margin minus DK", 18, CLAY, "Blank when either home margin is missing."),
    ("fd_minus_dk_note", "FD vs DK note", 36, CLAY, "Set when one of the two spreads is missing."),
    ("fd_novig_away_pct", "FD no-vig away %", 16, CLAY, "Book price after removing the hold. Not our probability."),
    ("fd_novig_home_pct", "FD no-vig home %", 16, CLAY, "Book price after removing the hold."),
    ("fd_hold_pts", "FD hold pts", 12, CLAY, "How far the two implied prices sit above 100."),
    ("dk_novig_away_pct", "DK no-vig away %", 16, CLAY, "Book price after removing the hold. Not our probability."),
    ("dk_novig_home_pct", "DK no-vig home %", 16, CLAY, "Book price after removing the hold."),
    ("dk_hold_pts", "DK hold pts", 12, CLAY, "Negative means the pair is not a two-way market. See the caution column."),
    ("novig_caution", "No-vig caution", 42, CLAY, "Present when a hold is negative."),
    ("fd_novig_note", "FD no-vig note", 22, CLAY, "missing-input when FanDuel had no moneyline pair."),
    ("dk_novig_note", "DK no-vig note", 22, CLAY, "missing-input when DraftKings had no moneyline pair."),
    ("props_flag", "Props", 36, CLAY, "Player props were not on these feeds. Inventory stays open."),
    ("identity", "Yardage identity", 42, CLAY, "Rush plus net pass versus the box-score total. Football only."),
    ("snapshot", "Retrieved (UTC)", 24, SLATE, "When this snapshot was pulled. Not a live feed."),
    ("sample_filter", "Sample filter", 42, SLATE, "Other seasons and season types skipped."),
]


def write_table(ws, columns, rows):
    header_font = Font(bold=True, color=WHITE, name="Calibri", size=11)
    header_align = Alignment(wrap_text=True, vertical="center")
    for col, (key, header, width, color, _help) in enumerate(columns, start=1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col)].width = width
    for row_index, row in enumerate(rows, start=2):
        for col, (key, _header, _width, _color, _help) in enumerate(columns, start=1):
            value = row.get(key)
            cell = ws.cell(row_index, col, value)
            cell.font = Font(name="Calibri", size=11)
            if isinstance(value, float):
                cell.number_format = "0.0"
            elif isinstance(value, int) and key.endswith(("price", "pct", "pts")) or (
                isinstance(value, int) and "price" in key
            ):
                cell.number_format = "+0;-0;0"
            elif isinstance(value, int):
                cell.number_format = "0"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.auto_filter.ref
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.sheet_properties.tabColor = columns[0][3]
    ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"


def write_start(ws, manual_text, games):
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    title = Font(bold=True, size=16, color=NAVY, name="Calibri")
    label = Font(bold=True, size=12, color=NAVY, name="Calibri")
    body = Font(size=11, name="Calibri")
    wrap = Alignment(wrap_text=True, vertical="top")
    ws["A1"] = "Sideline projections"
    ws["A1"].font = title
    ws.merge_cells("A1:B1")
    intro = (
        "One file for the October 30, 2026 snapshot. It holds the instructions, "
        "all four models, the 2,000-draw sim, and the FanDuel and DraftKings prices. "
        "Retrieved 2026-10-30 1:55 PM ET (17:55 UTC). This is not a live feed. "
        "Empty cells were blank in the snapshot. Nothing here was filled in afterward. "
        "Novig is not in this file."
    )
    ws["A3"] = "What this is"
    ws["B3"] = intro
    ws["A3"].font = label
    ws["B3"].font = body
    ws["B3"].alignment = wrap
    ws.row_dimensions[3].height = 48
    counts = []
    by_sport = {}
    for game in games:
        by_sport.setdefault(game["sport"], 0)
        by_sport[game["sport"]] += 1
    for sport, count in by_sport.items():
        counts.append(f"{sport}: {count}")
    ws["A5"] = "Games"
    ws["B5"] = f"{len(games)} games. " + "; ".join(counts) + ". MLB had no upcoming game, so it has no rows."
    ws["A5"].font = label
    ws["B5"].font = body
    ws["B5"].alignment = wrap
    ws.row_dimensions[5].height = 32
    sheets = (
        "Games is one row per game and is the sheet to filter, sort, or export. "
        "Methods is four rows per game, with each model minus each book. "
        "History is every for/against line: sample size, mean, median, mode, high, low. "
        "Football yards is rushing, gross passing, net passing, sack yards lost, and total offense. "
        "Open gaps is what this snapshot does not claim."
    )
    ws["A7"] = "Sheets"
    ws["B7"] = sheets
    ws["A7"].font = label
    ws["B7"].font = body
    ws["B7"].alignment = wrap
    ws.row_dimensions[7].height = 48
    rules = (
        "The four models are listed side by side. They are not averaged. "
        "Agreement is a vote count. It is not outside consensus. "
        "A line value is a book number that sits above every model or below every model. That is a point gap. "
        "Priced edge stays not labeled. The sim win rate is not a validated probability, and a gap versus a no-vig price is not an edge. "
        "Percent columns are already in percent: 42.9 means 42.9%, not 0.429. "
        "American odds use a plus or minus sign. "
        "Home spread is the board number: negative means the home team is favored. "
        "Home margin is how many points the home team is favored by, so it is the spread times -1. "
        "A price is ranked only when both books post the same number."
    )
    ws["A9"] = "How to read a row"
    ws["B9"] = rules
    ws["A9"].font = label
    ws["B9"].font = body
    ws["B9"].alignment = wrap
    ws.row_dimensions[9].height = 78
    ws["A11"] = "Rebuild"
    ws["B11"] = (
        "From the repo root: python3 scripts/export-slate-workbook.py. "
        "The script reads docs/projections/2026-10-30-slate.md and this manual. "
        "It does not call ESPN, FanDuel, or DraftKings."
    )
    ws["A11"].font = label
    ws["B11"].font = body
    ws["B11"].alignment = wrap
    ws.row_dimensions[11].height = 32

    ws["A13"] = "Games columns"
    ws["A13"].font = label
    row = 14
    for key, header, _width, _color, help_text in GAME_COLUMNS:
        ws.cell(row, 1, header).font = Font(bold=True, name="Calibri", size=11)
        cell = ws.cell(row, 2, help_text)
        cell.font = body
        cell.alignment = wrap
        ws.row_dimensions[row].height = 18
        row += 1

    row += 1
    ws.cell(row, 1, "Operating manual").font = label
    row += 1
    for paragraph in manual_text.splitlines():
        cell = ws.cell(row, 1, paragraph)
        cell.font = body
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell.alignment = wrap
        ws.row_dimensions[row].height = 18 if paragraph else 8
        row += 1
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.oddHeader.left.text = "Sideline projections — October 30, 2026 snapshot"


def write_commandments(ws, text):
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 110
    header = Font(bold=True, color=WHITE, name="Calibri", size=11)
    ws["A1"] = "Number"
    ws["B1"] = "Commandment"
    for cell in (ws["A1"], ws["B1"]):
        cell.font = header
        cell.fill = PatternFill("solid", fgColor=NAVY)
    row = 2
    for line in text.splitlines():
        match = re.fullmatch(r"(\d+)\. (.+)", line.strip())
        if not match:
            continue
        ws.cell(row, 1, int(match.group(1))).font = Font(bold=True, name="Calibri", size=12)
        body = ws.cell(row, 2, match.group(2))
        body.font = Font(name="Calibri", size=12)
        body.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 36
        row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{max(2, row - 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def write_open_gaps(ws):
    ws.sheet_properties.tabColor = CLAY
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 88
    header = Font(bold=True, color=WHITE, name="Calibri")
    ws["A1"] = "Gap"
    ws["B1"] = "Status in this snapshot"
    for cell in (ws["A1"], ws["B1"]):
        cell.font = header
        cell.fill = PatternFill("solid", fgColor=CLAY)
    gaps = [
        ("Player-stat inventory", "Open. Not recovered and not treated as canceled."),
        ("Player props", "FanDuel and DraftKings game-line feeds had no verified player prop."),
        ("FanDuel opening prices", "Not on the Ohio board."),
        ("DraftKings website", "Access denied. Open and latest prices are the ESPN quote."),
        ("FanDuel soccer", "No public Ohio board for the Premier League or MLS. Those games are DraftKings only."),
        ("Team totals", "Unavailable at FanDuel and DraftKings on these feeds."),
        ("Outside consensus", "Covers handicapper counts and the ESPN Matchup Predictor, in their own view. A missing row stays missing."),
        ("Our win probability", "Not calculated. A score is not turned into a probability."),
        ("Priced edge", "Not labeled. A gap versus the no-vig price is recorded and is not called an edge."),
        ("Novig", "Left out of this file. The pass compares FanDuel and DraftKings only."),
        ("MLB", "No upcoming game on the retrieved board."),
    ]
    for index, (gap, status) in enumerate(gaps, start=2):
        ws.cell(index, 1, gap).font = Font(name="Calibri", size=11)
        ws.cell(index, 2, status).font = Font(name="Calibri", size=11)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:B12"


def method_rows(game):
    rows = []
    for name in METHOD_ORDER:
        base = game["methods"].get(name, {})
        delta = game["comparison_methods"].get(name, {})
        rows.append(
            {
                "order": game["order"],
                "sport": game["sport"],
                "away": game["away"],
                "home": game["home"],
                "method": name,
                "away_score": base.get("away"),
                "home_score": base.get("home"),
                "home_margin": base.get("margin"),
                "combined": base.get("combined"),
                "minus_fd_margin": delta.get("minus_fd"),
                "minus_dk_margin": delta.get("minus_dk"),
                "minus_fd_total": delta.get("minus_fd_total"),
                "minus_dk_total": delta.get("minus_dk_total"),
            }
        )
    return rows


METHOD_COLUMNS = [
    ("order", "Order", 8, TEAL, ""),
    ("sport", "Sport", 28, TEAL, ""),
    ("away", "Away", 26, TEAL, ""),
    ("home", "Home", 26, TEAL, ""),
    ("method", "Method", 24, TEAL, ""),
    ("away_score", "Away score", 14, TEAL, ""),
    ("home_score", "Home score", 14, TEAL, ""),
    ("home_margin", "Home margin", 14, TEAL, ""),
    ("combined", "Combined", 12, TEAL, ""),
    ("minus_fd_margin", "Margin minus FD", 16, CLAY, ""),
    ("minus_dk_margin", "Margin minus DK", 16, CLAY, ""),
    ("minus_fd_total", "Total minus FD", 16, CLAY, ""),
    ("minus_dk_total", "Total minus DK", 16, CLAY, ""),
]

HISTORY_COLUMNS = [
    ("order", "Order", 8, GREEN, ""),
    ("sport", "Sport", 28, GREEN, ""),
    ("away", "Away", 26, GREEN, ""),
    ("home", "Home", 26, GREEN, ""),
    ("team", "Team", 26, GREEN, ""),
    ("stat", "Stat", 32, GREEN, ""),
    ("split", "Split", 12, GREEN, ""),
    ("n", "n", 8, GREEN, ""),
    ("mean", "Mean", 10, GREEN, ""),
    ("median", "Median", 10, GREEN, ""),
    ("mode", "Mode", 20, GREEN, ""),
    ("high", "High", 10, GREEN, ""),
    ("low", "Low", 10, GREEN, ""),
    ("note", "Note", 28, GREEN, ""),
]

YARD_COLUMNS = [
    ("order", "Order", 8, NAVY, ""),
    ("sport", "Sport", 28, NAVY, ""),
    ("away", "Away", 26, NAVY, ""),
    ("home", "Home", 26, NAVY, ""),
    ("stat", "Stat", 24, NAVY, ""),
    ("away_venue", "Away venue mean", 18, NAVY, ""),
    ("home_venue", "Home venue mean", 18, NAVY, ""),
    ("combined", "Combined", 12, NAVY, ""),
    ("away_overall", "Away overall mean", 18, NAVY, ""),
    ("home_overall", "Home overall mean", 18, NAVY, ""),
    ("away_for_n", "Away for n", 12, NAVY, ""),
    ("away_against_n", "Away against n", 14, NAVY, ""),
    ("home_for_n", "Home for n", 12, NAVY, ""),
    ("home_against_n", "Home against n", 14, NAVY, ""),
    ("away_note", "Away sample", 28, NAVY, ""),
    ("home_note", "Home sample", 28, NAVY, ""),
    ("identity", "Identity", 42, NAVY, ""),
]


def history_rows(game):
    rows = []
    for item in game["history"]:
        rows.append(
            {
                "order": game["order"],
                "sport": game["sport"],
                "away": game["away"],
                "home": game["home"],
                "team": item["team"],
                "stat": item["stat"],
                "split": item["split"],
                "n": item["n"],
                "mean": item["mean"],
                "median": item["median"],
                "mode": item["mode"],
                "high": item["high"],
                "low": item["low"],
                "note": item["note"],
            }
        )
    return rows


def yard_rows(game):
    if game.get("scoring_unit") != "points" or game["sport"].startswith("NBA"):
        return []
    if not game.get("identity") and all(
        game["yards"][key]["away_venue"] is None and game["yards"][key]["combined"] is None
        for key, _label in YARD_STATS
    ):
        return []
    rows = []
    for key, label in YARD_STATS:
        item = game["yards"][key]
        rows.append(
            {
                "order": game["order"],
                "sport": game["sport"],
                "away": game["away"],
                "home": game["home"],
                "stat": label,
                "away_venue": item["away_venue"],
                "home_venue": item["home_venue"],
                "combined": item["combined"],
                "away_overall": item["away_overall"],
                "home_overall": item["home_overall"],
                "away_for_n": item["away_for_n"],
                "away_against_n": item["away_against_n"],
                "home_for_n": item["home_for_n"],
                "home_against_n": item["home_against_n"],
                "away_note": item["away_note"],
                "home_note": item["home_note"],
                "identity": game.get("identity") if label == "Total offensive yards" else None,
            }
        )
    return rows


def validate(games, unparsed):
    errors = []
    if unparsed:
        errors.append(f"{len(unparsed)} unparsed bullets, first: {unparsed[0]}")
    if len(games) != 151:
        errors.append(f"expected 151 games, got {len(games)}")
    by_name = {game["matchup"]: game for game in games}
    ravens = by_name["Ravens at Bills"]
    expect(errors, ravens["methods"]["Venue-matched mean"]["away"], 26.2, "Ravens venue away")
    expect(errors, ravens["methods"]["Venue-matched mean"]["home"], 25, "Ravens venue home")
    expect(errors, ravens["methods"]["Venue-matched mean"]["margin"], -1.2, "Ravens venue margin")
    expect(errors, ravens["methods"]["Last 3 games"]["margin"], 3.5, "Ravens last3 margin")
    expect(errors, ravens.get("sim_away_median"), 28.5, "Ravens sim away")
    expect(errors, ravens.get("sim_home_median"), 24.5, "Ravens sim home")
    expect(errors, ravens.get("sim_home_win_pct"), 42.9, "Bills sim win")
    expect(errors, ravens.get("sim_away_win_pct"), 55.9, "Ravens sim win")
    expect(errors, ravens.get("sim_tie_pct"), 1.3, "Ravens sim tie")
    expect(errors, ravens.get("fd_book_margin"), 3, "FD Bills margin")
    expect(errors, ravens.get("dk_book_margin"), 3, "DK Bills margin")
    expect(errors, ravens.get("dk_novig_away"), 39.1, "DK no-vig Ravens")
    expect(errors, ravens.get("dk_novig_home"), 60.9, "DK no-vig Bills")
    expect(errors, ravens["market"]["Side"]["dk_open"], "-2.5 -110", "DK open side")
    expect(errors, ravens.get("agreement"), "split", "Ravens agreement")
    expect(errors, ravens.get("visitor_synergy"), "diverge", "Ravens visitor synergy")
    expect(errors, ravens.get("home_synergy"), "reinforce", "Ravens home synergy")
    points = next(item for item in ravens["history"] if item["team"] == "Ravens" and item["stat"] == "points for" and item["split"] == "overall")
    expect(errors, points["n"], 7, "Ravens n")
    expect(errors, points["mean"], 24.1, "Ravens mean")
    expect(errors, points["mode"], "no repeated value", "Ravens mode")
    expect(errors, ravens["yards"]["rushing yards"]["away_venue"], 139.5, "Ravens rush")

    titans = by_name["Titans at Bengals"]
    expect(errors, titans.get("fd_side_class"), "outside-method-range", "Titans FD side")
    expect(errors, titans.get("dk_side_class"), "outside-method-range", "Titans DK side")

    hornets = by_name["Hornets at Wizards"]
    expect(errors, hornets.get("sim_sample", "").startswith("thin"), True, "Hornets thin sim")
    expect(errors, hornets.get("fd_book_total"), 219.5, "Hornets FD total")
    expect(errors, hornets.get("dk_book_total"), 218.5, "Hornets DK total")
    expect(errors, hornets.get("fd_total_class"), "outside-method-range", "Hornets FD total class")

    united = by_name["Man Utd at Chelsea"]
    expect(errors, united["methods"]["Venue-matched mean"]["away"], 1.5, "United goals")
    expect(errors, united["methods"]["Venue-matched mean"]["home"], 2.1, "Chelsea goals")
    expect(errors, united.get("fd_side_class"), "missing-input", "United FD side")
    expect(errors, united.get("dk_total_class"), "outside-method-range", "Chelsea DK total")
    expect(errors, united.get("dk_hold") < 0, True, "Chelsea negative hold")

    dc = by_name["D.C. United at Columbus"]
    expect(errors, dc.get("home_synergy"), "flat", "DC home synergy")

    neutral = [game["matchup"] for game in games if game.get("site") and game["site"].startswith("neutral")]
    if len(neutral) != 4:
        errors.append(f"expected 4 neutral sites, got {neutral}")

    missing_methods = [game["matchup"] for game in games if len(game["methods"]) != 4]
    if missing_methods:
        errors.append(f"games missing a method: {missing_methods[:5]}")
    missing_sim = [game["matchup"] for game in games if game.get("sim_away_median") is None]
    if missing_sim:
        errors.append(f"games missing a sim median: {missing_sim[:5]}")
    return errors


def expect(errors, actual, wanted, label):
    if actual != wanted:
        errors.append(f"{label}: expected {wanted!r}, got {actual!r}")


def build(slate_path, manual_path, out_path):
    text = slate_path.read_text()
    manual = manual_path.read_text()
    games, unparsed, _retrieved = parse_games(text)
    errors = validate(games, unparsed)
    if errors:
        raise SystemExit("\n".join(errors))

    wb = Workbook()
    commandments = wb.active
    commandments.title = "Commandments"
    write_commandments(commandments, DEFAULT_COMMANDMENTS.read_text())
    start = wb.create_sheet("Start here")
    write_start(start, manual, games)

    game_sheet = wb.create_sheet("Games")
    game_rows = [game_row(game) for game in games]
    write_table(game_sheet, GAME_COLUMNS, game_rows)
    csv_path = out_path.with_name("sideline-games-2026-10-30.csv")
    write_games_csv(csv_path, game_rows)

    methods = wb.create_sheet("Methods")
    method_data = []
    for game in games:
        method_data.extend(method_rows(game))
    write_table(methods, METHOD_COLUMNS, method_data)

    history = wb.create_sheet("History")
    history_data = []
    for game in games:
        history_data.extend(history_rows(game))
    write_table(history, HISTORY_COLUMNS, history_data)

    yards = wb.create_sheet("Football yards")
    yard_data = []
    for game in games:
        yard_data.extend(yard_rows(game))
    write_table(yards, YARD_COLUMNS, yard_data)

    gaps = wb.create_sheet("Open gaps")
    write_open_gaps(gaps)

    wb.properties.title = "Sideline all-sports projections 2026-10-30"
    wb.properties.subject = "Four models, resampling sim, FanDuel and DraftKings. Snapshot, not a live feed."
    wb.properties.creator = "Sideline"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {
        "games": len(games),
        "methods": len(method_data),
        "history": len(history_data),
        "yards": len(yard_data),
        "path": out_path,
        "csv": csv_path,
    }


def write_games_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([header for _key, header, _width, _color, _help in GAME_COLUMNS])
        for row in rows:
            writer.writerow(
                ["" if row.get(key) is None else row.get(key) for key, _header, _width, _color, _help in GAME_COLUMNS]
            )


def main():
    slate = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SLATE
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    summary = build(slate, DEFAULT_MANUAL, out)
    print(
        f"wrote {summary['path']} and {summary['csv']} games={summary['games']} "
        f"methods={summary['methods']} history={summary['history']} yards={summary['yards']}"
    )


if __name__ == "__main__":
    main()
